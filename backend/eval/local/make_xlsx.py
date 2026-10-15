import sys
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

out = sys.argv[1]
wb = Workbook()
ws = wb.active
ws.title = "요약"
F = "맑은 고딕"
thin = Side(style="thin", color="BFBFBF")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
head_fill = PatternFill("solid", fgColor="E8EEF7")
yellow = PatternFill("solid", fgColor="FFF2CC")
grey = PatternFill("solid", fgColor="F3F4F6")
wrap = Alignment(wrap_text=True, vertical="top")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)

ws["B1"] = "네오러다이트 AI 3차 평가 — 요약"; ws["B1"].font = Font(name=F, size=14, bold=True)
ws["B2"] = "노란 칸만 적어 주시면 됩니다. 답변 원문은 함께 보내 드린 Word 문서 5장에 있습니다."
ws["B2"].font = Font(name=F, size=10, color="6B7280")

cols = ["문항", "지난번 점수\n(정확성/형식/근거)", "지난번 대비", "정확성 (1~5)", "답변 형식 (1~5)", "근거 법령 (1~5)", "한 줄 의견"]
widths = [30, 16, 12, 11, 11, 11, 60]
for i, (c, w) in enumerate(zip(cols, widths)):
    cell = ws.cell(row=4, column=2 + i, value=c)
    cell.font = Font(name=F, bold=True, size=10); cell.fill = head_fill; cell.border = box; cell.alignment = center
    ws.column_dimensions[chr(ord("B") + i)].width = w
ws.column_dimensions["A"].width = 2
ws.row_dimensions[4].height = 32

items = [
    ("예시", "—", "개선", 4, 3, 5, "세율표가 숫자로 나와 좋아졌으나 설명 문장이 길다"),
    ("질문1. 종합소득세율", "3 / 2 / 4"), ("질문2. 연봉 5천만원 세액 계산", "4 / 3 / 4"),
    ("질문3. 제조업 간이과세", "2 / 2 / 2"), ("질문4. 법인세 익금", "3 / 2 / 2"),
    ("질문1~4 한 대화 연속", "—"),
    ("추가1. 법인세율", "—"), ("추가2. 증여세 세율 구간", "—"), ("추가3. 연봉 7천만원 실수령액", "—"),
]
dv_score = DataValidation(type="whole", operator="between", formula1="1", formula2="5", allow_blank=True,
                          error="1~5 사이 숫자를 넣어 주세요", errorTitle="점수")
dv_cmp = DataValidation(type="list", formula1='"개선,약간 개선,비슷함,퇴보"', allow_blank=True)
ws.add_data_validation(dv_score); ws.add_data_validation(dv_cmp)
r = 5
for it in items:
    vals = list(it) + [None] * (7 - len(it))
    example = it[0] == "예시"
    for j, v in enumerate(vals):
        c = ws.cell(row=r, column=2 + j, value=v)
        c.font = Font(name=F, size=10, color="6B7280" if example else "000000", italic=example)
        c.border = box
        c.alignment = wrap if j in (0, 6) else center
        if j >= 2 and not example:
            c.fill = yellow
        elif j == 1:
            c.fill = grey
    if not example:
        dv_cmp.add(ws.cell(row=r, column=4))
        for j in (5, 6, 7):
            dv_score.add(ws.cell(row=r, column=j))
    ws.row_dimensions[r].height = 42
    r += 1

r += 1
notes = ["점수 기준",
         "정확성: 세법·계산·요건 판단이 맞는가 (5 = 세무사 답변과 같음)",
         "답변 형식: 표·번호 목록·계산 단계처럼 읽기 쉬운 형식인가",
         "근거 법령: 관련 조문을 맞게 가져오고, 법령 → 예규·판례 순서로 근거를 댔는가",
         "지난번 점수는 10월 재평가 때 주신 점수입니다(질문1~4). 추가 질문·연속 대화는 이번에 처음 점수를 받습니다."]
for i, t in enumerate(notes):
    c = ws.cell(row=r + i, column=2, value=t)
    c.font = Font(name=F, size=10, bold=(i == 0), color="000000" if i == 0 else "374151")
ws.freeze_panes = "B5"

ws2 = wb.create_sheet("문장별 의견")
ws2["B1"] = "문장별 의견 (선택)"; ws2["B1"].font = Font(name=F, size=14, bold=True)
ws2["B2"] = "짚어 주실 문장을 Word 문서에서 복사해 붙이고 의견을 적어 주세요."
ws2["B2"].font = Font(name=F, size=10, color="6B7280")
h2 = ["문항", "문장 (복사)", "구분", "의견"]
w2 = [24, 60, 14, 50]
dv_kind = DataValidation(type="list", formula1='"법적 오류,문장·문법,형식,제안,기타"', allow_blank=True)
ws2.add_data_validation(dv_kind)
for i, (c, w) in enumerate(zip(h2, w2)):
    cell = ws2.cell(row=4, column=2 + i, value=c)
    cell.font = Font(name=F, bold=True, size=10); cell.fill = head_fill; cell.border = box; cell.alignment = center
    ws2.column_dimensions[chr(ord("B") + i)].width = w
ws2.column_dimensions["A"].width = 2
for rr in range(5, 35):
    for j in range(4):
        c = ws2.cell(row=rr, column=2 + j); c.border = box; c.alignment = wrap; c.font = Font(name=F, size=10)
        c.fill = yellow
    dv_kind.add(ws2.cell(row=rr, column=4))
    ws2.row_dimensions[rr].height = 36
ws2.freeze_panes = "B5"
wb.save(out)
print("written", out)
